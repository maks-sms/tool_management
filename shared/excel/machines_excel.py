# shared/excel/machines_excel.py
"""
Экспорт инструмента на станках в Excel (.xlsx)
с объединением ячеек по станку. Без UI.
"""
import sqlite3
from collections import OrderedDict
from typing import List

from shared.config import get_db_path, LOCATION_MACHINE_PREFIX


def fetch_machines_for_export() -> "OrderedDict[str, list]":
    """
    Возвращает OrderedDict:
        { 'станок': [ (код, артикул, тип, диаметр, длина, бренд, кол-во, дата), ... ] }
    """
    conn = sqlite3.connect(get_db_path(), timeout=10)
    cur = conn.cursor()
    cur.execute("""
        SELECT p.note          AS machine,
               t.inner_code,
               t.articul,
               t.type,
               t.diameter,
               t.length,
               t.brand,
               p.qty,
               p.date
        FROM placement p
        JOIN tools t ON t.id = p.tool_id
        WHERE p.location LIKE ?
        ORDER BY p.note, t.inner_code
    """, (f"{LOCATION_MACHINE_PREFIX}%",))
    rows = cur.fetchall()
    conn.close()

    grouped = OrderedDict()
    for machine, code, art, ttype, diam, length, brand, qty, date in rows:
        machine = (machine or "—").strip()
        grouped.setdefault(machine, []).append(
            (code or "", art or "", ttype or "", diam or "",
             length or "", brand or "", qty or 0, date or "")
        )
    return grouped


def export_machines_to_excel(file_path: str) -> int:
    """
    Экспортирует инструмент на станках в .xlsx.

    :param file_path: путь к файлу
    :return: количество выгруженных строк
    """
    groups = fetch_machines_for_export()
    if not groups:
        return 0

    _write_excel(file_path, groups)
    total = sum(len(items) for items in groups.values())
    return total


def _write_excel(file_path: str, groups) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Инструмент на станках"

    headers = ["№", "Станок", "Код", "Артикул", "Тип",
               "Диаметр", "Длина", "Производитель", "Количество"]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_align = Alignment(horizontal="center", vertical="center",
                             wrap_text=True)

    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    machine_align = Alignment(horizontal="center", vertical="center",
                              wrap_text=True)

    row = 2
    counter = 1

    for machine, items in groups.items():
        start_row = row

        for (code, art, ttype, diam, length, brand, qty, date) in items:
            ws.cell(row=row, column=1, value=counter).alignment = center
            ws.cell(row=row, column=3, value=code).alignment = center
            ws.cell(row=row, column=4, value=art).alignment = left
            ws.cell(row=row, column=5, value=ttype).alignment = center
            ws.cell(row=row, column=6, value=diam).alignment = center
            ws.cell(row=row, column=7, value=length).alignment = center
            ws.cell(row=row, column=8, value=brand).alignment = left
            ws.cell(row=row, column=9, value=qty).alignment = center

            for c in range(1, 10):
                ws.cell(row=row, column=c).border = border

            row += 1
            counter += 1

        end_row = row - 1

        machine_cell = ws.cell(row=start_row, column=2, value=machine)
        machine_cell.alignment = machine_align
        machine_cell.font = Font(bold=True, size=11)
        machine_cell.border = border

        if end_row > start_row:
            ws.merge_cells(
                start_row=start_row, start_column=2,
                end_row=end_row, end_column=2
            )
            for r in range(start_row, end_row + 1):
                ws.cell(row=r, column=2).border = border

    widths = {1: 6, 2: 22, 3: 12, 4: 30, 5: 18,
              6: 10, 7: 10, 8: 18, 9: 12}
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{row - 1}"

    wb.save(file_path)