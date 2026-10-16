# shared/excel/tools_excel.py
"""
Экспорт склада (инструментов) в Excel (.xlsx).
Без UI — просто принимает путь к файлу и пишет.
"""
import sqlite3
from typing import List, Tuple

from shared.config import get_db_path


def fetch_tools_for_export(only_available: bool = False) -> List[Tuple]:
    """
    Возвращает список кортежей:
        (inner_code, articul, type, diameter, length, brand, analog, qty_total)
    """
    conn = sqlite3.connect(get_db_path(), timeout=10)
    cur = conn.cursor()

    if only_available:
        cur.execute("""SELECT inner_code, articul, type, diameter, length,
                              brand, analog, qty_total
                       FROM tools
                       WHERE qty_total > 0
                       ORDER BY inner_code""")
    else:
        cur.execute("""SELECT inner_code, articul, type, diameter, length,
                              brand, analog, qty_total
                       FROM tools ORDER BY inner_code""")

    rows = cur.fetchall()
    conn.close()
    return rows


def export_tools_to_excel(file_path: str,
                          only_available: bool = False) -> int:
    """
    Экспортирует инструменты склада в .xlsx.

    :param file_path: путь, куда сохранить файл
    :param only_available: True — только остаток > 0
    :return: количество выгруженных строк
    """
    rows = fetch_tools_for_export(only_available)
    if not rows:
        return 0

    _write_excel(file_path, rows, only_available)
    return len(rows)


def _write_excel(file_path: str, rows, only_available: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Склад"

    headers = ["№", "Код", "Артикул", "Остаток", "Тип",
               "Диаметр", "Длина", "Аналоги", "Производитель"]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_align = Alignment(horizontal="center", vertical="center",
                             wrap_text=True)

    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    zero_fill = PatternFill("solid", fgColor="FDECEA")

    for i, r in enumerate(rows, start=1):
        code, art, t, d, l, b, a, q = r
        row_idx = i + 1

        ws.cell(row=row_idx, column=1, value=i).alignment = center
        ws.cell(row=row_idx, column=2, value=code or "").alignment = center
        ws.cell(row=row_idx, column=3, value=art or "").alignment = left
        ws.cell(row=row_idx, column=4, value=q or 0).alignment = center
        ws.cell(row=row_idx, column=5, value=t or "").alignment = center
        ws.cell(row=row_idx, column=6, value=d or "").alignment = center
        ws.cell(row=row_idx, column=7, value=l or "").alignment = center
        ws.cell(row=row_idx, column=8, value=a or "").alignment = left
        ws.cell(row=row_idx, column=9, value=b or "").alignment = left

        for c in range(1, 10):
            cell = ws.cell(row=row_idx, column=c)
            cell.border = border
            if not only_available and (q or 0) == 0:
                cell.fill = zero_fill

    widths = {1: 6, 2: 12, 3: 32, 4: 10, 5: 18,
              6: 10, 7: 10, 8: 28, 9: 20}
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(rows) + 1}"

    wb.save(file_path)